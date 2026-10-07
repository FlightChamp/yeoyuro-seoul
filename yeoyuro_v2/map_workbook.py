"""
yeoyuro_v2/map_workbook.py
==========================
노선도 좌표 기본 파일을 **읽기 전용**으로 해석한다. 이 모듈은 사용자 파일에 절대 쓰지 않는다.

원칙
----
- 사용자가 고치는 값: Station_Visual_Nodes 의 x_px, y_px  /  Station_Labels 의 label_dx, label_dy
                      /  Station_Click_Areas 의 click_dx, click_dy  (이동량, 기본 0)
- 계산 값: 역 클릭 위치 = 그 역 노선별 점 평균 + click 이동량, 역명 위치 = 점 평균 + label 이동량.
  엑셀 안에서는 같은 계산을 수식으로 보여 주지만, 이 모듈은 수식 캐시를 믿지 않고 직접 계산한다
  (엑셀을 연 채로 / 저장 직후 / 스크립트로 좌표를 바꾼 직후 모두 같은 결과).
"""

from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

import pandas as pd

W, H = 5120, 2880
SHEETS = ["Station_Display_Master", "Station_Visual_Nodes", "Station_Click_Areas", "Station_Labels",
          "Line_Sequences", "Transfer_Links"]


def read_all(path: Path) -> dict[str, pd.DataFrame]:
    """엑셀이 파일을 열고 있어도 읽을 수 있게 임시 사본에서 읽는다."""
    tmp = Path(tempfile.mkdtemp()) / "wb.xlsx"
    shutil.copy(path, tmp)
    xl = pd.ExcelFile(tmp)
    out = {s: xl.parse(s) for s in xl.sheet_names if not s.startswith("_")}
    xl.close()
    return out


def materialize(path: Path) -> dict:
    """파일에서 읽어 계산한다. (열린 엑셀에서 읽은 표는 materialize_frames 에 바로 넘긴다)"""
    s = read_all(path)
    s["_literal"] = literal_cells(path)          # 수식 대신 숫자를 직접 입력한 역명·클릭 칸
    return materialize_frames(s)


def signature(s: dict) -> int:
    """미리보기에 영향을 주는 입력값(점 좌표·이동량)이 바뀌었는지 판단하는 지문."""
    parts = []
    for name, cols in (("Station_Visual_Nodes", ["visual_node_id", "x_px", "y_px"]),
                       ("Station_Labels", ["station_key", "label_dx", "label_dy"]),
                       ("Station_Click_Areas", ["station_key", "click_dx", "click_dy"])):
        df = s.get(name)
        if df is not None:
            parts.append(df[[c for c in cols if c in df.columns]].astype(str).to_csv(index=False))
    return hash("".join(parts))


def materialize_frames(s: dict) -> dict:
    """사용자 입력(점 좌표 + 이동량)으로 클릭·역명 실제 좌표를 계산한 시트들과 점 겹침 목록을 돌려준다."""
    s = dict(s)
    vn = s["Station_Visual_Nodes"].copy()
    vn["x_px"] = pd.to_numeric(vn.x_px, errors="coerce")
    vn["y_px"] = pd.to_numeric(vn.y_px, errors="coerce")
    mean = vn.groupby("station_key")[["x_px", "y_px"]].mean()
    vn["x_pct"], vn["y_pct"] = vn.x_px / W * 100, vn.y_px / H * 100

    lit = s.pop("_literal", {})

    def place(df, xc, yc, dxc, dyc, sheet):
        """좌표 = 직접 입력한 숫자가 있으면 그 값, 수식(또는 빈칸)이면 역 점 평균 + 이동량."""
        df = df.copy()
        dx = pd.to_numeric(df[dxc], errors="coerce").fillna(0).values if dxc in df else 0
        dy = pd.to_numeric(df[dyc], errors="coerce").fillna(0).values if dyc in df else 0
        m = mean.reindex(df.station_key)
        auto_x, auto_y = (m.x_px.values + dx).round(), (m.y_px.values + dy).round()
        typed = lit.get(sheet, {})
        tx = [typed.get((k, xc)) for k in df.station_key]
        ty = [typed.get((k, yc)) for k in df.station_key]
        df[xc] = [t if t is not None else a for t, a in zip(tx, auto_x)]
        df[yc] = [t if t is not None else a for t, a in zip(ty, auto_y)]
        df["_source"] = ["직접 입력" if (a is not None or b is not None) else "자동" for a, b in zip(tx, ty)]
        df[xc.replace("_px", "_pct")] = df[xc] / W * 100
        df[yc.replace("_px", "_pct")] = df[yc] / H * 100
        return df

    ca = place(s["Station_Click_Areas"], "click_x_px", "click_y_px", "click_dx", "click_dy", "Station_Click_Areas")
    lb = place(s["Station_Labels"], "label_x_px", "label_y_px", "label_dx", "label_dy", "Station_Labels")
    g = vn.dropna(subset=["x_px", "y_px"]).groupby(["x_px", "y_px"]).station_key.agg(lambda x: sorted(set(x)))
    clash = [(int(x), int(y), ks) for (x, y), ks in g.items() if len(ks) > 1]
    out = dict(s)
    out["_literal_used"] = lit
    out.update({"Station_Visual_Nodes": vn, "Station_Click_Areas": ca, "Station_Labels": lb, "_clash": clash,
                "_missing": vn[vn.x_px.isna() | vn.y_px.isna()].visual_node_id.tolist()})
    return out


def literal_cells(path: Path) -> dict:
    """역명·클릭 시트에서 수식이 아니라 숫자가 직접 입력된 좌표 칸 {(역, 컬럼): 값}."""
    import openpyxl
    tmp = Path(tempfile.mkdtemp()) / "f.xlsx"
    shutil.copy(path, tmp)
    wb = openpyxl.load_workbook(tmp, read_only=True)
    out = {}
    for sheet, cols in (("Station_Labels", ("label_x_px", "label_y_px")),
                        ("Station_Click_Areas", ("click_x_px", "click_y_px"))):
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        head = list(next(rows))
        ik = head.index("station_key")
        idx = {c: head.index(c) for c in cols if c in head}
        d = {}
        for r in rows:
            k = r[ik]
            if k is None:
                continue
            for c, i in idx.items():
                v = r[i]
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    d[(k, c)] = float(v)
        out[sheet] = d
    wb.close()
    return out


def snapshot(m: dict) -> dict:
    """변경 비교용: 최종 좌표(final) + 사용자가 고칠 수 있는 입력 칸(inputs)."""
    vn, lb, ca = m["Station_Visual_Nodes"], m["Station_Labels"], m["Station_Click_Areas"]
    final = {f"{r.visual_node_id} 점": (r.x_px, r.y_px) for r in vn.itertuples()}
    final.update({f"{r.station_key} 역명": (r.label_x_px, r.label_y_px) for r in lb.itertuples()})
    final.update({f"{r.station_key} 클릭": (r.click_x_px, r.click_y_px) for r in ca.itertuples()})
    lit = m.get("_literal_used", {})
    inputs = {}
    for r in vn.itertuples():
        inputs[("점", r.visual_node_id, "x_px")] = r.x_px
        inputs[("점", r.visual_node_id, "y_px")] = r.y_px
    for df, kind, pre, sheet in ((lb, "역명", "label", "Station_Labels"), (ca, "클릭", "click", "Station_Click_Areas")):
        for r in df.itertuples():
            k = r.station_key
            for ax in ("dx", "dy"):
                inputs[(kind, k, f"{pre}_{ax}")] = getattr(r, f"{pre}_{ax}", 0)
            for ax in ("x", "y"):
                inputs[(kind, k, f"{pre}_{ax}_px")] = lit.get(sheet, {}).get((k, f"{pre}_{ax}_px"), "자동")
    return {"final": final, "inputs": inputs}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def diff(old: dict | None, new: dict) -> list[str]:
    """고친 칸마다: 무엇을 → 반영 여부 → 반영 안 됐으면 이유·해결법. 이어서 따라 움직인 것."""
    if old is None:
        return []
    of, nf, oi, ni = old["final"], new["final"], old["inputs"], new["inputs"]

    def moved(key):
        a, b = of.get(key), nf.get(key)
        if a is None or b is None:
            return a != b
        return abs((_num(a[0]) or 0) - (_num(b[0]) or 0)) > 0.4 or abs((_num(a[1]) or 0) - (_num(b[1]) or 0)) > 0.4

    out, explained = [], set()
    for key in sorted(set(oi) | set(ni), key=str):
        a, b = oi.get(key), ni.get(key)
        if a == b or (_num(a) is not None and _num(b) is not None and abs(_num(a) - _num(b)) < 1e-9):
            continue
        kind, k, col = key
        target = f"{k} 점" if kind == "점" else f"{k} {kind}"
        if kind == "점":
            target = f"{k} 점"
        shown = lambda v: "자동(수식)" if v == "자동" else (f"{_num(v):.0f}" if _num(v) is not None else "빈칸")
        line = f"{kind} {k}: {col} {shown(a)} → {shown(b)}"
        if moved(target):
            out.append(line + "  ✔ 반영됨")
            explained.add(target)
        elif col.endswith(("_dx", "_dy")):
            ax = col[-1]
            fixed = ni.get((kind, k, col.replace(f"_d{ax}", f"_{ax}_px")))
            if fixed != "자동":
                out.append(line + f"  ✘ 반영 안 됨 — {col.replace(f'_d{ax}', f'_{ax}_px')} 칸에 숫자 {shown(fixed)} 이(가) 직접 입력돼 있어 "
                           f"이동량이 무시됩니다. 그 칸의 숫자를 지우면(빈칸) 자동 + 이동량으로 바뀝니다")
            else:
                out.append(line + "  ✘ 반영 안 됨 (원인 불명 — 이 줄을 그대로 알려 주세요)")
        else:
            out.append(line + "  (위치 변화 없음)")
    for key in nf:
        if key not in explained and moved(key) and not key.endswith(" 점"):
            a, b = of.get(key), nf.get(key)
            out.append(f"  ↳ {key} 따라 이동: ({_num(a[0]):.0f}, {_num(a[1]):.0f}) → ({_num(b[0]):.0f}, {_num(b[1]):.0f})")
    return out


def write_values_copy(m: dict, out: Path) -> Path:
    """계산된 값으로 가져오기 스크립트(14)용 사본을 만든다 (수식 없음)."""
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        for name, df in m.items():
            if not name.startswith("_"):
                df.drop(columns=[c for c in df.columns if str(c).startswith("_")]).to_excel(
                    w, sheet_name=name[:31], index=False)
    return out
