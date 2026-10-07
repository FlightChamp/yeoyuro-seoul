"""
tests/test_line9.py
===================
v2.6 9호선 mart·엔진 통합 검사. 원본 없이 커밋된 data/marts/v2/line9 만으로 돈다.

1. 그래프 구조: 9L 38역(양방향 74 edge), 9X 16역(30 edge), 환승 9개 연결 × 패턴, 일반↔급행 16역
2. 엔진: 급행 김포공항→종합운동장 평일 08:00 = 시각표 실측 46.4분 (bin 별 운행시간 반영)
3. 1~8호선만 쓰는 OD 는 9호선 추가 전과 같은 결과
4. 이벤트 9호선 제외 옵션, 급행 미운행 시간대 차단
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
M9 = ROOT / "data" / "marts" / "v2" / "line9"


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    if not (M9 / "route_edges.parquet").exists():
        pytest.skip("9호선 mart 없음")
    from yeoyuro_v2 import load_v1
    from yeoyuro_v2 import line9 as L9
    from yeoyuro_v2.time_dependent import TDRouter
    mod, sr, disp = load_v1(ROOT)
    r9 = L9.build_shadow_root(ROOT, tmp_path_factory.mktemp("v26") / "base", "base")
    d9 = pd.read_csv(r9 / "data" / "master" / "station_display_master.csv")
    return {"mod": mod, "sr": sr, "disp": disp, "r9": r9, "d9": d9,
            "rt9": TDRouter(r9, mod, sr, d9, "weekday"), "rt": TDRouter(ROOT, mod, sr, disp, "weekday")}


def test_graph_structure():
    if not (M9 / "route_edges.parquet").exists():
        pytest.skip("9호선 mart 없음")
    e = pd.read_parquet(M9 / "route_edges.parquet")
    assert (e.line_id == "9L").sum() == 74 and (e.line_id == "9X").sum() == 30
    t = pd.read_parquet(M9 / "transfer_tf181.parquet")
    inter = t[~(t.from_line.isin(["9L", "9X"]) & t.to_line.isin(["9L", "9X"]))]
    assert set(inter.station_name) == {"김포공항", "당산", "여의도", "동작", "고속터미널", "종합운동장", "석촌", "올림픽공원"}
    sw = t[t.from_line.isin(["9L", "9X"]) & t.to_line.isin(["9L", "9X"])]
    assert sw.station_name.nunique() == 16 and (sw.transfer_time_min == 0).all()
    f100 = pd.read_parquet(M9 / "transfer_tf100.parquet").set_index(["from_node", "to_node"]).transfer_time_min
    f181 = t.set_index(["from_node", "to_node"]).transfer_time_min
    nz = f100[f100 > 0]
    assert ((f181.loc[nz.index] / nz).round(2) == 1.81).all()


def test_express_matches_timetable(env):
    from yeoyuro_v2.time_dependent import parse_hhmm
    res = env["rt9"].evaluate_od("김포공항", "종합운동장", parse_hhmm("08:00"))
    best = res[0]
    assert {n.split("_")[0] for n in best.path} == {"9X"}
    assert best.actual_time_min == pytest.approx(46.4, abs=0.3)


def test_lines_1_8_unchanged(env):
    from yeoyuro_v2.time_dependent import parse_hhmm
    a = env["rt"].evaluate_od("노원", "수유", parse_hhmm("08:00"))[0]
    b = env["rt9"].evaluate_od("노원", "수유", parse_hhmm("08:00"))[0]
    assert a.path == b.path and a.actual_time_min == pytest.approx(b.actual_time_min)


def test_event_exclude_and_service(env):
    from yeoyuro_v2.time_dependent import TimeDependentEvaluator, parse_hhmm
    mod, sr = env["mod"], env["sr"]
    rs = mod.RouteScorer(env["r9"], "saturday", "18:00", "2025-09-27")        # 불꽃축제 (여의도 포함)
    hw = sr.load_headway(env["r9"])
    path = ["9L_여의도", "9L_샛강", "9L_노량진"]
    on = TimeDependentEvaluator(rs, hw, "step").evaluate(path, parse_hhmm("18:00"))
    off = TimeDependentEvaluator(rs, hw, "step", event_exclude_lines=("9L", "9X")).evaluate(path, parse_hhmm("18:00"))
    assert on.max_congestion > off.max_congestion
    rs2 = mod.RouteScorer(env["r9"], "weekday", "05:30")
    ev2 = TimeDependentEvaluator(rs2, sr.load_headway(env["r9"]), "step")
    early = ev2.evaluate(["9X_김포공항", "9X_마곡나루"], parse_hhmm("05:05"))
    assert early.out_of_window                       # 05:05 에는 급행 출발이 없다


def test_cap160_scaling(tmp_path):
    if not (M9 / "route_edges.parquet").exists():
        pytest.skip("9호선 mart 없음")
    from yeoyuro_v2 import line9 as L9
    r = L9.build_shadow_root(ROOT, tmp_path / "cap", "cap160")
    lk = pd.read_parquet(r / "data/marts/congestion_edge_lookup.parquet")
    med = pd.read_parquet(M9 / "congestion_median.parquet")
    a = lk[lk.line_id.isin(["9L", "9X"])].set_index(["station_uid", "direction", "day_type", "time_bin_index"]).congestion_median
    b = med.set_index(["station_uid", "direction", "day_type", "time_bin_index"]).congestion_median
    j = pd.concat([a.rename("cap"), b.rename("med")], axis=1).dropna()
    assert len(j) > 1000
    assert ((j.cap - j.med * L9.CAP160_FACTOR).abs() <= 0.006).all()   # 소수 둘째 자리 반올림 오차 이내
    one8 = lk[~lk.line_id.isin(["9L", "9X"])]
    assert len(one8) == len(pd.read_parquet(ROOT / "data/marts/congestion_edge_lookup.parquet"))


def test_single_s23_module_for_spawn():
    """28 이 27·23b 를 함께 불러와도 s23 모듈은 하나여야 한다 (Windows spawn 피클링 오류 방지)."""
    import importlib.util
    for k in ("s23", "s27", "s23b", "s28"):
        sys.modules.pop(k, None)
    spec = importlib.util.spec_from_file_location("s28", ROOT / "scripts" / "v2" / "28_line9_capacity_check.py")
    m = importlib.util.module_from_spec(spec)
    sys.modules["s28"] = m
    spec.loader.exec_module(m)
    s23 = sys.modules["s23"]
    assert m.S27.S23 is s23 and m.S23B.S23 is s23
    import pickle
    assert pickle.loads(pickle.dumps(s23._init_worker)) is s23._init_worker


def test_base_map_workbook_has_lines_1_to_9():
    """v2.6 새 기본 좌표 워크북: 1~9호선, 가져오기 스크립트가 읽는 6개 시트, 좌표 빈칸 없음."""
    p = ROOT / "data" / "master" / "yeoyuro_seoul_vector_map_coordinate_workbook.xlsx"
    xl = pd.ExcelFile(p)
    for s in ("Station_Display_Master", "Station_Visual_Nodes", "Station_Click_Areas", "Station_Labels",
              "Line_Sequences", "Transfer_Links"):
        assert s in xl.sheet_names
    v = xl.parse("Station_Visual_Nodes")
    assert sorted(v.line_id.unique().tolist()) == list(range(1, 10))
    assert len(v) == 315 and v.x_px.notna().all() and v.y_px.notna().all()
    assert (v.line_id == 9).sum() == 38
    seq = xl.parse("Line_Sequences")
    assert seq[seq.line_id == 9].path_order.tolist() == list(range(1, 39))


def test_live_workbook_follows_nodes_and_keeps_label_offsets(tmp_path):
    """수식 연동 구조: 역 점을 옮기면 역명·클릭이 따라오고, 역명 이동량은 유지되며, 검사는 파일을 바꾸지 않는다."""
    import importlib.util
    import shutil
    import openpyxl
    from yeoyuro_v2.map_workbook import materialize
    p = tmp_path / "wb.xlsx"
    shutil.copy(ROOT / "data" / "master" / "yeoyuro_seoul_vector_map_coordinate_workbook.xlsx", p)
    spec = importlib.util.spec_from_file_location("mig29", ROOT / "scripts" / "v2" / "29_migrate_live_formulas.py")
    mig = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mig)
    assert mig.migrate(p) == 0 and mig.migrate(p) == 0               # 두 번째는 '이미 변환됨'
    wb = openpyxl.load_workbook(p)
    ws, lb = wb["Station_Visual_Nodes"], wb["Station_Labels"]
    hv = {c.value: c.column for c in ws[1] if c.value}
    hl = {c.value: c.column for c in lb[1] if c.value}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, hv["visual_node_id"]).value == "흑석_L9":
            ws.cell(r, hv["x_px"], 2129)
    for r in range(2, lb.max_row + 1):
        if lb.cell(r, hl["station_key"]).value == "흑석":
            lb.cell(r, hl["label_dy"], 50)
    wb.save(p)
    before = p.stat().st_mtime
    m = materialize(p)
    L = m["Station_Labels"].set_index("station_key")
    C = m["Station_Click_Areas"].set_index("station_key")
    assert (L.at["흑석", "label_x_px"], L.at["흑석", "label_y_px"]) == (2129, 1800)
    assert (C.at["흑석", "click_x_px"], C.at["흑석", "click_y_px"]) == (2129, 1750)
    assert all("흑석" not in ks for _, _, ks in m["_clash"])
    assert p.stat().st_mtime == before                                 # 읽기 전용
