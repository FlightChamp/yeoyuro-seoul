"""
scripts/v2/29_build_base_map_workbook.py
========================================
1~9호선 노선도 좌표 워크북을 **처음부터 새로** 만든다 (v2.6 새 기본 파일, source of truth).

입력
----
    --base   : 기존 1~8호선 워크북 (v1 기본 파일)
    --line9  : 9호선 좌표를 입력한 워크북 (사용자 입력본, Station_Visual_Nodes 의 line_id=9 행을 읽음)
출력
----
    data/master/yeoyuro_seoul_vector_map_coordinate_workbook.xlsx   (새 기본 파일)

정리 원칙
---------
- 시트 이름은 v1 가져오기 스크립트(14)가 읽는 6개를 그대로 쓴다.
- Station_Visual_Nodes / Line_Sequences 는 1→9호선, 노선별 path_order 순으로 정렬.
- 모든 좌표·퍼센트는 수식이 아니라 숫자 값.
- 사용자 입력 좌표는 그대로 쓴다. 다른 역과 좌표가 완전히 같으면 '확인 필요'로 표시만 한다 (고치지 않음).
- 오래된 안내(4900×2757 기준)·AI_Handoff_Prompt·9호선 입력도우미 시트는 넣지 않는다.

실행
----
    python scripts/v2/29_build_base_map_workbook.py --base <1~8호선 워크북> --line9 <9호선 입력 워크북>
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "master" / "yeoyuro_seoul_vector_map_coordinate_workbook.xlsx"
W, H = 5120, 2880
LINES = [  # line_id, 이름, 색, 범위, 지선, 비고
    (1, "1호선", "#0052A4", "서울역~청량리", "main", "서울교통공사 구간"),
    (2, "2호선", "#009D3E", "순환선 전 구간 + 성수지선 + 신정지선", "main; seongsu_branch; sinjeong_branch", "순환선은 렌더링 시 마지막→처음 연결"),
    (3, "3호선", "#EF7C1C", "지축~오금", "main", ""),
    (4, "4호선", "#00A5DE", "불암산~남태령", "main", ""),
    (5, "5호선", "#996CAC", "방화~하남검단산, 강동~마천", "common; hanam_branch; macheon_branch", "강동에서 분기"),
    (6, "6호선", "#CD7C2F", "응암~신내 (응암순환 포함)", "main", "응암순환은 단방향"),
    (7, "7호선", "#747F00", "장암~온수", "main", ""),
    (8, "8호선", "#E6186C", "암사역사공원~모란", "main", ""),
    (9, "9호선", "#BDB092", "개화~중앙보훈병원", "main", "노선도는 한 줄. 일반/급행 구분은 경로 계산(9L/9X)에서만 사용"),
]
TRANSFERS9 = {"김포공항": ["5"], "당산": ["2"], "여의도": ["5"], "동작": ["4"], "고속터미널": ["3", "7"],
              "종합운동장": ["2"], "석촌": ["8"], "올림픽공원": ["5"]}
HEAD = PatternFill("solid", fgColor="DDEBF7")
WARN = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(*(Side(style="thin", color="D9D9D9"),) * 4)


def light(hex_color: str, a: float = 0.78) -> str:
    """노선 색을 밝게 (셀 배경용)."""
    h = hex_color.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{int(c + (255 - c) * a):02X}" for c in (r, g, b))


def write_sheet(wb, name, df: pd.DataFrame, widths=None, line_col=None, warn_rows=()):
    ws = wb.create_sheet(name)
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(1, j, c)
        cell.font, cell.fill, cell.border = Font(bold=True), HEAD, THIN
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    colors = {str(l[0]): l[2] for l in LINES}
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(i, j, v)
            c.border = THIN
        if line_col is not None:
            lid = str(df.iloc[i - 2][line_col])
            if lid in colors:
                ws.cell(i, df.columns.get_loc(line_col) + 1).fill = PatternFill("solid", fgColor=light(colors[lid]))
        if (i - 2) in warn_rows:
            for j in range(1, len(df.columns) + 1):
                if j != (df.columns.get_loc(line_col) + 1 if line_col else -1):
                    ws.cell(i, j).fill = WARN
    for j, c in enumerate(df.columns, start=1):
        longest = df[c].dropna().astype(str).str.len().max() if len(df) else 0
        longest = 0 if pd.isna(longest) else longest
        w = (widths or {}).get(c, max(10, min(36, int(longest * 1.3))))
        ws.column_dimensions[get_column_letter(j)].width = max(w, len(str(c)) * 1.6)
    ws.freeze_panes = "B2" if name != "안내" else None
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    return ws


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--line9", required=True)
    ap.add_argument("--out", default=str(OUT))
    args = ap.parse_args(argv)
    b = {s: pd.read_excel(args.base, sheet_name=s) for s in
         ["Config", "Station_Display_Master", "Station_Visual_Nodes", "Station_Click_Areas", "Station_Labels",
          "Line_Sequences", "Transfer_Links", "Map_Layer_Order", "Sources"]}
    u = pd.read_excel(args.line9, sheet_name="Station_Visual_Nodes")
    useq = pd.read_excel(args.line9, sheet_name="Line_Sequences")

    # ---------------- Station_Visual_Nodes: 1~8 (사용자 수정 반영) + 9
    vn = b["Station_Visual_Nodes"].copy()
    upd = u.set_index("visual_node_id")[["x_px", "y_px"]]
    changed = []
    for i, r in vn.iterrows():
        if r.visual_node_id in upd.index:
            nx, ny = upd.loc[r.visual_node_id]
            if (nx, ny) != (r.x_px, r.y_px):
                changed.append((r.visual_node_id, int(r.x_px), int(r.y_px), int(nx), int(ny)))
                vn.at[i, "x_px"], vn.at[i, "y_px"] = int(nx), int(ny)
    n9 = u[u.line_id.astype(str) == "9"].copy()
    n9["mark_status"] = "완료"
    n9["memo"] = ""
    vn = pd.concat([vn, n9[vn.columns]], ignore_index=True)
    seq = pd.concat([b["Line_Sequences"], useq[useq.line_id.astype(str) == "9"]], ignore_index=True)
    br_order = {"main": 0, "common": 0, "seongsu_branch": 1, "sinjeong_branch": 2, "hanam_branch": 1, "macheon_branch": 2}
    seq["_b"] = seq.branch_code.map(br_order).fillna(9)
    seq = seq.sort_values(["line_id", "_b", "path_order"]).drop(columns="_b").reset_index(drop=True)
    first = seq.drop_duplicates("visual_node_id").reset_index().set_index("visual_node_id")["index"]
    vn["_o"] = vn.visual_node_id.map(first)
    vn = vn.sort_values(["line_id", "_o"]).drop(columns="_o").reset_index(drop=True)
    vn["visual_node_no"] = range(1, len(vn) + 1)
    vn["x_px"], vn["y_px"] = vn.x_px.astype(int), vn.y_px.astype(int)
    vn["x_pct"], vn["y_pct"] = (vn.x_px / W * 100).round(4), (vn.y_px / H * 100).round(4)
    # 다른 역끼리 좌표가 같은 경우 → 확인 필요 (값은 고치지 않음)
    grp = vn.groupby(["x_px", "y_px"]).station_key.transform(lambda s: s.nunique())
    clash = vn.index[grp > 1].tolist()
    for i in clash:
        others = sorted(set(vn[(vn.x_px == vn.at[i, "x_px"]) & (vn.y_px == vn.at[i, "y_px"])].station_key) - {vn.at[i, "station_key"]})
        vn.at[i, "mark_status"] = "확인 필요"
        vn.at[i, "memo"] = f"다른 역({', '.join(others)})과 좌표가 같음 — 확인 필요"

    # ---------------- Station_Display_Master
    dm = b["Station_Display_Master"].copy()
    for i, r in dm.iterrows():
        if r.station_key in TRANSFERS9:
            dm.at[i, "available_lines"] = f"{r.available_lines},9"
            dm.at[i, "project_branches"] = f"{r.project_branches};9:main"
            dm.at[i, "is_transfer_station"] = True
    new9 = [s for s in seq[seq.line_id == 9].station_key if s not in set(dm.station_key)]
    add = pd.DataFrame([{"station_no": len(dm) + k + 1, "station_key": s, "station_name": s, "display_name": s + "역",
                         "available_lines": "9", "candidate_station_uids": None, "project_branches": "9:main",
                         "is_transfer_station": False, "is_event_station": False, "alias_or_note": None,
                         "memo": None} for k, s in enumerate(new9)])
    dm = pd.concat([dm, add], ignore_index=True)

    # ---------------- Click / Labels (기존 240 그대로 + 9호선 전용역 30)
    ca, lb = b["Station_Click_Areas"].copy(), b["Station_Labels"].copy()
    node = vn.set_index("visual_node_id")
    for s in new9:
        x, y = int(node.loc[f"{s}_L9", "x_px"]), int(node.loc[f"{s}_L9", "y_px"])
        ca.loc[len(ca)] = {"mark_status": "완료", "station_key": s, "station_name": s, "display_name": s + "역",
                           "available_lines": "9", "click_x_px": x, "click_y_px": y, "click_radius_px": 24,
                           "click_shape": "circle", "source_method": "auto_from_visual_nodes", "manual_override_note": None}
        lb.loc[len(lb)] = {"mark_status": "완료", "station_key": s, "station_name": s, "display_name": s + "역",
                           "available_lines": "9", "label_x_px": x, "label_y_px": y, "label_size_px": 10.5,
                           "label_weight": "normal", "label_anchor": "center", "label_rotation": 0,
                           "source_method": "default_from_visual_nodes", "memo": None}
    # 기존 환승역의 노선 목록 갱신 + 1~8 노드가 바뀐 역은 클릭·라벨도 같이 이동
    moved = {c[0].rsplit("_L", 1)[0]: c for c in changed}
    for t in (ca, lb):
        for i, r in t.iterrows():
            if r.station_key in TRANSFERS9 and "9" not in str(r.available_lines).split(","):
                t.at[i, "available_lines"] = f"{r.available_lines},9"
    for s, (_, ox, oy, nx, ny) in moved.items():
        n_nodes = (vn.station_key == s).sum()
        if n_nodes == 1:
            for t, xc, yc in ((ca, "click_x_px", "click_y_px"), (lb, "label_x_px", "label_y_px")):
                i = t.index[t.station_key == s][0]
                t.at[i, xc] += nx - ox
                t.at[i, yc] += ny - oy
    for t, xc, yc in ((ca, "click_x_px", "click_y_px"), (lb, "label_x_px", "label_y_px")):
        t[xc.replace("_px", "_pct")] = (t[xc] / W * 100).round(4)
        t[yc.replace("_px", "_pct")] = (t[yc] / H * 100).round(4)

    # ---------------- Transfer_Links
    tl = b["Transfer_Links"].copy()
    for s, lines in TRANSFERS9.items():
        for ln in lines:
            tl.loc[len(tl)] = {"station_key": s, "station_name": s, "from_line_id": int(ln), "to_line_id": 9,
                               "from_visual_node_id": f"{s}_L{ln}", "to_visual_node_id": f"{s}_L9",
                               "transfer_style": "connector", "connector_visible": True, "connector_color": "#9CA3AF",
                               "memo": "지도 렌더링용 환승 연결선"}
    tl = tl.sort_values(["from_line_id", "to_line_id", "station_key"]).reset_index(drop=True)

    # ---------------- 새 워크북
    wb = Workbook()
    wb.remove(wb.active)
    guide = pd.DataFrame({"항목": [
        "이 파일", "좌표계", "핵심 원칙", "Lines", "Station_Display_Master", "Station_Visual_Nodes", "Line_Sequences",
        "Transfer_Links", "Station_Click_Areas", "Station_Labels", "수정 방법", "검사", "앱 반영", "확인 필요 표시"],
        "설명": [
        "여유로 서울 1~9호선 벡터 노선도 좌표의 기본 파일(source of truth). 2026-10 v2.6 에서 처음부터 다시 정리.",
        f"캔버스 {W} × {H} px, 원점 왼쪽 위, y 는 아래로 증가. pct 는 px 에서 계산한 값(수식 아님).",
        "사용자 선택 = station_key / 지도 그리기 = visual_node_id / 경로 계산 = station_uid (9호선은 9L·9X). 세 계층을 섞지 않는다.",
        "노선 목록·색·범위. 색을 바꾸려면 여기와 가져오기 스크립트의 LINE_COLORS 를 같이 바꾼다.",
        "역 단위 목록 270개. 사용자가 고르는 역. candidate_station_uids 는 가져오기 스크립트가 채운다.",
        "노선별 역 점 좌표 315개 (1→9호선, 노선 순서). 환승역은 노선마다 점이 따로 있고, 같은 역의 점은 겹쳐도 된다.",
        "노선을 그리는 순서. visual_node_id 를 path_order 대로 잇는다.",
        "환승역 안에서 노선별 점을 잇는 회색 연결선 (보이는 선일 뿐, 환승 시간 계산과 무관).",
        "클릭 판정용 투명 원. 역당 1개.",
        "역명 위치. 역당 1개. 겹치면 label_x_px / label_y_px 만 고친다.",
        "좌표를 고치면 x_px, y_px 만 바꾼다 (pct 는 가져오기 때 다시 계산됨). 새 역을 넣으면 Visual_Nodes·Line_Sequences·Display_Master·Click·Labels 에 같이 넣는다.",
        "python scripts/v2/29_check_map_workbook.py → 검증표 + 미리보기 (앱 노선도는 바뀌지 않음)",
        "Release 단계에서 9호선 경로 그래프와 함께 반영 (python scripts/14_import_map_workbook.py).",
        "mark_status = '확인 필요' 인 행은 다른 역과 좌표가 완전히 같은 경우. 값은 사용자 입력 그대로 두었다."]})
    ws = write_sheet(wb, "안내", guide, widths={"항목": 24, "설명": 120})
    for r in range(2, len(guide) + 2):
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    cfg = b["Config"]
    cfg = cfg[cfg.parameter != "line9_color"].copy()
    write_sheet(wb, "Config", cfg, widths={"parameter": 40, "value": 26, "note": 60})
    lines = pd.DataFrame(LINES, columns=["line_id", "line_name", "color", "project_range", "branch_codes", "note"])
    ws = write_sheet(wb, "Lines", lines, widths={"project_range": 40, "branch_codes": 40, "note": 60})
    for i, l in enumerate(LINES, start=2):
        ws.cell(i, 3).fill = PatternFill("solid", fgColor=l[2].lstrip("#"))
        ws.cell(i, 3).font = Font(color="FFFFFF", bold=True)
    write_sheet(wb, "Station_Display_Master", dm, line_col=None)
    write_sheet(wb, "Station_Visual_Nodes", vn, line_col="line_id", warn_rows=set(clash),
                widths={"memo": 44, "alias_or_note": 18})
    write_sheet(wb, "Line_Sequences", seq, line_col="line_id")
    write_sheet(wb, "Transfer_Links", tl, line_col="to_line_id")
    write_sheet(wb, "Station_Click_Areas", ca)
    write_sheet(wb, "Station_Labels", lb)
    write_sheet(wb, "Map_Layer_Order", b["Map_Layer_Order"], widths={"note": 60, "source_sheet": 36})
    src = b["Sources"]
    src = src[~src.source_name.isin(["coordinate_image_basis", "workbook_basis"])]
    src = pd.concat([src, pd.DataFrame([
        {"source_name": "서울메트로9호선_역사정보", "url": "철도데이터포털(KRIC) 서울시메트로9호선주식회사 제공 (2026-05-31)",
         "usage_note": "9호선 38역 역번호·위경도·승강장 유형. 역 순서 검증"},
        {"source_name": "v1 좌표 워크북 (1~8호선)", "url": "data/master/archive/ (보관)",
         "usage_note": "1~8호선 좌표의 원본. 방화·개화산 좌표는 v2.6 에서 사용자가 조정"},
        {"source_name": "9호선 좌표", "url": "사용자 직접 입력 (2026-10)", "usage_note": "9호선 38역 x_px, y_px"}])],
        ignore_index=True)
    write_sheet(wb, "Sources", src, widths={"url": 60, "usage_note": 60})
    notes = pd.DataFrame(
        [{"구분": "확인 필요", "내용": vn.at[i, "visual_node_id"] + " — " + vn.at[i, "memo"]} for i in clash] +
        [{"구분": "1~8호선 좌표 변경", "내용": f"{c[0]}: ({c[1]}, {c[2]}) → ({c[3]}, {c[4]})"} for c in changed] +
        [{"구분": "정리", "내용": "AI_Handoff_Prompt·Project_Scope(→ Lines)·9호선 입력도우미 시트 제거, 안내 갱신 (5120×2880 기준)"}])
    write_sheet(wb, "변경_기록", notes, widths={"구분": 18, "내용": 100})
    wb.save(args.out)
    print(f"저장: {args.out}")
    print(f"- 노선 9개 · 역 {len(dm)}개 · 노드 {len(vn)}개 · 노선 순서 {len(seq)}행 · 환승 연결선 {len(tl)}개")
    print(f"- 1~8호선 좌표 변경 {len(changed)}개 반영: " + ", ".join(c[0] for c in changed))
    print(f"- 확인 필요 {len(clash)}개: " + ", ".join(vn.loc[clash, 'visual_node_id']))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
