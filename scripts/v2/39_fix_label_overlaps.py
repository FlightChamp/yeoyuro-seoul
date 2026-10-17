"""
scripts/v2/39_fix_label_overlaps.py
===================================
노선도 역명 겹침을 찾아, 겹치는 역명만 위·좌·우 등 빈 자리로 옮기는 label_dx / label_dy 를 계산한다.

    python scripts/v2/39_fix_label_overlaps.py            # 미리보기: 겹치는 쌍과 옮길 역명 목록만 출력 (파일 안 바꿈)
    python scripts/v2/39_fix_label_overlaps.py --apply    # 기본 좌표 엑셀의 Station_Labels label_dx/label_dy 에 기록

- 역 점(Station_Visual_Nodes)·노선 순서·경로 계산은 건드리지 않는다. 역명 이동량 두 칸만 바꾼다.
- 겹치지 않는 역명, 사용자가 이미 이동량을 넣은 역명은 그대로 둔다 (--include-manual 로 포함 가능).
- 크기 계산 기준: 1920×1080 발표 화면에서 앱 노선도 1 캔버스 단위 ≈ 0.30 화면 px (경로 추천·네트워크 지도 공통 근사).
  한글 1글자 ≈ 글자 크기 1em. 확대하면 글자 간격은 넓어지므로, 기본 화면에서 안 겹치면 확대해도 안 겹친다.
- 적용 뒤 앱을 열면 엑셀이 노선도보다 새것이라 자동으로 다시 가져온다. 결과가 마음에 안 들면 git 으로 되돌린다:
    git checkout -- data/master/yeoyuro_seoul_vector_map_coordinate_workbook.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from yeoyuro_v2.map_workbook import materialize          # noqa: E402
from yeoyuro_v2.proc import utf8_stdout                   # noqa: E402

BASE = ROOT / "data" / "master" / "yeoyuro_seoul_vector_map_coordinate_workbook.xlsx"
SCALE = 0.30                    # 화면 px / 캔버스 단위 (1920×1080, 앱 노선도 기본 화면)
FONT_SCALE = 0.95               # 앱 LABEL_FONT_SCALE 과 같게
DY_NORMAL, DY_MAJOR = 32, 44    # 앱 LABEL_DY_NORMAL / LABEL_DY_MAJOR 와 같게
R_SMALL, R_TRANSFER = 3.5 / SCALE + 2, 6.5 / SCALE + 3
GAP = 3 / SCALE                 # 글자와 점·다른 글자 사이 최소 여백 (화면 3px)


def text_w(s: str, size: float) -> float:
    em = size * FONT_SCALE
    w = 0.0
    for ch in str(s):
        if "가" <= ch <= "힣":
            w += 1.0
        elif ch.isdigit() or ch.isascii():
            w += 0.6
        else:
            w += 0.45
    return w * em * 1.04 / SCALE           # 굵은 글씨 약간 넓게


def text_h(size: float) -> float:
    return size * FONT_SCALE * 1.25 / SCALE


def load(path: Path):
    s = materialize(path)
    lab = s["Station_Labels"].copy()
    vn = s["Station_Visual_Nodes"].copy()
    seq = s["Line_Sequences"].copy()
    for c in ("label_x_px", "label_y_px", "label_size_px", "label_dx", "label_dy"):
        lab[c] = pd.to_numeric(lab[c], errors="coerce").fillna(0.0)
    for c in ("x_px", "y_px"):
        vn[c] = pd.to_numeric(vn[c], errors="coerce")
    vn["is_transfer_station"] = vn["is_transfer_station"].astype(str).str.lower().isin(["true", "1"])
    pos = vn.set_index("visual_node_id")[["x_px", "y_px"]]
    segs = []
    for (_, _), g in seq.sort_values(["line_id", "branch_code", "path_order"]).groupby(["line_id", "branch_code"]):
        ids = [i for i in g.visual_node_id if i in pos.index]
        for a, b in zip(ids, ids[1:]):
            segs.append((*pos.loc[a].values, *pos.loc[b].values))
    return lab, vn, segs


def base_boxes(lab: pd.DataFrame) -> dict:
    """역명마다 이동량 0 일 때의 기준점(역 점 평균)과 글자 크기."""
    out = {}
    for r in lab.itertuples():
        size = float(r.label_size_px) or 10.5
        out[r.station_key] = {"x": r.label_x_px - r.label_dx, "y": r.label_y_px - r.label_dy, "size": size,
                              "w": text_w(r.station_key, size), "h": text_h(size),
                              "dy0": DY_MAJOR if size > 10.5 else DY_NORMAL,
                              "dx": float(r.label_dx), "dyv": float(r.label_dy)}
    return out


def box(b: dict, dx: float, dy: float):
    cx, top = b["x"] + dx, b["y"] + dy + b["dy0"]
    return (cx - b["w"] / 2, top, cx + b["w"] / 2, top + b["h"])


def overlap(a, b, pad=0.0) -> float:
    w = min(a[2], b[2]) + pad - max(a[0], b[0])
    h = min(a[3], b[3]) + pad - max(a[1], b[1])
    return w * h if (w > 0 and h > 0) else 0.0


def seg_hits(bx, segs) -> int:
    """글자 상자를 지나는 노선 선분 수 (근사: 선분을 10등분한 점이 상자 안에 들어가는지)."""
    n = 0
    for x0, y0, x1, y1 in segs:
        if max(x0, x1) < bx[0] or min(x0, x1) > bx[2] or max(y0, y1) < bx[1] or min(y0, y1) > bx[3]:
            continue
        if any(bx[0] <= x0 + (x1 - x0) * t / 10 <= bx[2] and bx[1] <= y0 + (y1 - y0) * t / 10 <= bx[3]
               for t in range(11)):
            n += 1
    return n


def candidates(b: dict, r_own: float):
    w, h, d0 = b["w"], b["h"], b["dy0"]
    side_dy = -h / 2 - d0
    yield "아래(기본)", 0.0, 0.0
    yield "아래·오른쪽", w * 0.45, 0.0
    yield "아래·왼쪽", -w * 0.45, 0.0
    yield "아래·조금 오른쪽", w * 0.25, 0.0
    yield "아래·조금 왼쪽", -w * 0.25, 0.0
    yield "위", 0.0, -(r_own + GAP + h + d0)
    yield "위·조금 오른쪽", w * 0.25, -(r_own + GAP + h + d0)
    yield "위·왼쪽", -(r_own + GAP + w / 2), -(r_own + GAP + h + d0)
    yield "위·조금 왼쪽", -w * 0.25, -(r_own + GAP + h + d0)
    yield "오른쪽", r_own + GAP + w / 2, side_dy
    yield "왼쪽", -(r_own + GAP + w / 2), side_dy
    yield "조금 더 아래", 0.0, h * 0.8


# 자동 배치가 '다른 역 이름처럼' 읽히게 만든 곳은 위치를 고정한다 (시연 역인 동대문역사문화공원 주변, 2026-10-17 화면 확인)
PINNED = {"동대문역사문화공원": "위", "을지로4가": "아래·조금 오른쪽", "을지로3가": "아래·조금 왼쪽",
          "동대문": "위"}


def solve(lab, vn, segs, include_manual=False, passes=6):
    B = base_boxes(lab)
    markers = [(r.x_px, r.y_px, R_TRANSFER if r.is_transfer_station else R_SMALL, r.station_key) for r in vn.itertuples()]
    r_own = {k: max([m[2] for m in markers if m[3] == k] or [R_SMALL]) for k in B}
    cur = {k: (b["dx"], b["dyv"]) for k, b in B.items()}
    fixed = {k for k, b in B.items() if (b["dx"] or b["dyv"]) and not include_manual}


    def cost(k, dx, dy, mode):
        bx = box(B[k], dx, dy)
        c = 0.0
        for j, (ox, oy) in cur.items():
            if j != k:
                c += 40 * overlap(bx, box(B[j], ox, oy), GAP) / (B[k]["h"] ** 2)
        for mx, my, mr, mk in markers:
            c += 25 * overlap(bx, (mx - mr, my - mr, mx + mr, my + mr), GAP) / (B[k]["h"] ** 2)
        c += 2.0 * seg_hits(bx, segs)                         # 노선 선 위에 글자를 올리지 않는다
        c += 0.0 if mode.startswith("아래(기본)") else 0.8       # 기본 위치를 선호
        c += 0.006 * (dx * dx + dy * dy) ** 0.5               # 자기 역에서 멀어질수록 불리 (어느 역 이름인지 헷갈리지 않게)
        return c

    def conflicts():
        ks = list(B)
        pairs = []
        for i, a in enumerate(ks):
            ba = box(B[a], *cur[a])
            for b in ks[i + 1:]:
                if overlap(ba, box(B[b], *cur[b]), 0) > 0:
                    pairs.append((a, b))
        mk = []
        for a in ks:
            ba = box(B[a], *cur[a])
            for mx, my, mr, k in markers:
                if k != a and overlap(ba, (mx - mr, my - mr, mx + mr, my + mr), 0) > 0:
                    mk.append((a, k))
                    break
        return pairs, mk

    before = conflicts()
    for k, mode in PINNED.items():
        if k in B and k not in fixed:
            dx, dy = next((x, y) for m, x, y in candidates(B[k], r_own[k]) if m == mode)
            cur[k] = (round(dx), round(dy))
            fixed.add(k)
    for _ in range(passes):
        pairs, mk = conflicts()
        # 역명끼리 겹친 것만 옮긴다 (다른 역 점을 살짝 덮는 것까지 고치면 거의 모든 역명이 움직인다). 새 자리를 고를 때는 점도 피한다.
        involved = {a for p in pairs for a in p} - fixed
        if not involved:
            break
        changed = False
        for k in sorted(involved, key=lambda x: (-B[x]["w"], x)):     # 긴 역명부터 (같으면 이름순 → 실행마다 같은 결과)
            best = min(((cost(k, dx, dy, m), m, dx, dy) for m, dx, dy in candidates(B[k], r_own[k])),
                       key=lambda z: z[0])
            if (round(best[2]), round(best[3])) != (round(cur[k][0]), round(cur[k][1])):
                cur[k] = (round(best[2]), round(best[3]))
                changed = True
        if not changed:
            break
    after = conflicts()
    moves = [(k, B[k]["dx"], B[k]["dyv"], cur[k][0], cur[k][1]) for k in B
             if (round(cur[k][0]), round(cur[k][1])) != (round(B[k]["dx"]), round(B[k]["dyv"]))]
    return before, after, moves


def apply(moves, path: Path) -> int:
    import openpyxl
    if (path.parent / ("~$" + path.name)).exists():
        print("[중단] 엑셀이 기본 파일을 열고 있습니다. 엑셀을 닫고 다시 실행하세요.")
        return 3
    wb = openpyxl.load_workbook(path)
    ws = wb["Station_Labels"]
    h = {c.value: c.column for c in ws[1] if c.value}
    rows = {ws.cell(r, h["station_key"]).value: r for r in range(2, ws.max_row + 1)}
    for k, _, _, dx, dy in moves:
        ws.cell(rows[k], h["label_dx"]).value = int(dx)
        ws.cell(rows[k], h["label_dy"]).value = int(dy)
    wb.save(path)
    return 0


def main() -> int:
    utf8_stdout()
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--include-manual", action="store_true")
    ap.add_argument("--xlsx", default=str(BASE))
    a = ap.parse_args()
    path = Path(a.xlsx)
    lab, vn, segs = load(path)
    before, after, moves = solve(lab, vn, segs, a.include_manual)
    print("=" * 64)
    print(f"읽은 파일: {path}")
    print(f"역명끼리 겹침: {len(before[0])}쌍 → {len(after[0])}쌍 · 역명이 다른 역 점을 덮음: {len(before[1])}곳 → {len(after[1])}곳")
    for a_, b_ in before[0]:
        print(f"  겹침(전): {a_} ↔ {b_}")
    for a_, b_ in after[0]:
        print(f"  남은 겹침: {a_} ↔ {b_}")
    print(f"옮길 역명 {len(moves)}개 (label_dx, label_dy):")
    for k, dx0, dy0, dx, dy in moves:
        print(f"  {k:<12} ({dx0:.0f}, {dy0:.0f}) → ({dx:.0f}, {dy:.0f})")
    if a.apply and moves:
        rc = apply(moves, path)
        if rc:
            return rc
        print("엑셀에 기록했습니다. 앱을 열면 노선도가 자동으로 다시 그려집니다.")
    elif moves:
        print("미리보기만 했습니다. 기록하려면 --apply 를 붙여 다시 실행하세요.")
    print("=" * 64)
    return 0


if __name__ == "__main__":
    sys.exit(main())
