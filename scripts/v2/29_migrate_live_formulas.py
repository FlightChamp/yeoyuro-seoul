"""
scripts/v2/29_migrate_live_formulas.py
======================================
기본 좌표 파일을 "엑셀에서 고치면 바로 따라오는" 구조로 한 번 바꾼다. (실행 전 자동 백업)

바뀌는 것
---------
- Station_Labels      : label_dx, label_dy (이동량, 입력 칸) 추가. label_x_px/y_px/pct 는 수식
                        = 그 역 노선별 점 평균 + 이동량
- Station_Click_Areas : click_dx, click_dy 추가. click_x_px/y_px/pct 는 수식 (같은 방식)
- Station_Visual_Nodes: x_pct/y_pct 수식, 다른 역과 좌표가 같은 행은 조건부 서식으로 주황 표시 (실시간)
- 숨김 시트 _sync_state 삭제 (이전 동기화 방식 폐기)

지금 파일에 이미 직접 옮겨 둔 역명·클릭 위치는 "점 평균과의 차이"를 이동량으로 옮겨 그대로 보존한다.
이미 변환된 파일이면 아무것도 하지 않는다.

실행
----
    python scripts/v2/29_migrate_live_formulas.py
"""

from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "data" / "master" / "yeoyuro_seoul_vector_map_coordinate_workbook.xlsx"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="EDEDED")
ORANGE = PatternFill("solid", fgColor="FCE4D6")


def main(argv=None) -> int:
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(BASE))
    a = ap.parse_args(argv)
    return migrate(Path(a.xlsx))


def migrate(BASE: Path) -> int:
    if (BASE.parent / ("~$" + BASE.name)).exists():
        print("[중단] 엑셀이 기본 파일을 열고 있습니다. 이번 한 번만 엑셀을 닫고 실행하세요.")
        return 3
    wb = openpyxl.load_workbook(BASE)
    lab = wb["Station_Labels"]
    if any(c.value == "label_dx" for c in lab[1]):
        print("이미 변환된 파일입니다. 아무것도 바꾸지 않았습니다.")
        return 0
    bak = BASE.parent / "archive" / f"{BASE.stem}_before_live_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    bak.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(BASE, bak)

    vn = wb["Station_Visual_Nodes"]
    hv = {c.value: c.column for c in vn[1] if c.value}
    L = get_column_letter
    nodes = pd.DataFrame([(vn.cell(r, hv["station_key"]).value, vn.cell(r, hv["x_px"]).value,
                           vn.cell(r, hv["y_px"]).value) for r in range(2, vn.max_row + 1)],
                         columns=["k", "x", "y"]).dropna(subset=["k"])
    mean = nodes.groupby("k")[["x", "y"]].mean()
    K, X, Y = L(hv["station_key"]), L(hv["x_px"]), L(hv["y_px"])
    rng = lambda col: f"Station_Visual_Nodes!${col}$2:${col}$2000"

    # 역 점: pct 수식 + 겹침 실시간 표시
    for r in range(2, vn.max_row + 1):
        if vn.cell(r, hv["station_key"]).value is None:
            continue
        vn.cell(r, hv["x_pct"], f"={X}{r}/5120*100").fill = GREY
        vn.cell(r, hv["y_pct"], f"={Y}{r}/2880*100").fill = GREY
        for c in (hv["x_px"], hv["y_px"]):
            vn.cell(r, c).fill = YELLOW
        if vn.cell(r, hv["mark_status"]).value == "확인 필요":
            vn.cell(r, hv["mark_status"], "완료")
            vn.cell(r, hv["memo"], None)
            for c in range(1, vn.max_column + 1):     # 예전 '확인 필요' 주황 칠 지우기 (이제 조건부 서식이 대신함)
                if c not in (hv["line_id"], hv["x_px"], hv["y_px"], hv["x_pct"], hv["y_pct"]):
                    vn.cell(r, c).fill = PatternFill(fill_type=None)
    last = vn.max_row
    vn.conditional_formatting.add(
        f"A2:{L(vn.max_column)}{last}",
        FormulaRule(formula=[f'COUNTIFS(${X}$2:${X}${last},${X}2,${Y}$2:${Y}${last},${Y}2,${K}$2:${K}${last},"<>"&${K}2)>0'],
                    fill=ORANGE, stopIfTrue=False))

    # 역명·클릭: 이동량 칸 + 수식
    for ws, pre in ((lab, "label"), (wb["Station_Click_Areas"], "click")):
        h = {c.value: c.column for c in ws[1] if c.value}
        cdx, cdy = ws.max_column + 1, ws.max_column + 2
        for c, name in ((cdx, f"{pre}_dx"), (cdy, f"{pre}_dy")):
            hc = ws.cell(1, c, name)
            hc.font, hc.fill = Font(bold=True), YELLOW
            ws.column_dimensions[L(c)].width = 11
        SK = L(h["station_key"])
        for r in range(2, ws.max_row + 1):
            k = ws.cell(r, h["station_key"]).value
            if k is None:
                continue
            x0, y0 = ws.cell(r, h[f"{pre}_x_px"]).value, ws.cell(r, h[f"{pre}_y_px"]).value
            mx, my = (mean.at[k, "x"], mean.at[k, "y"]) if k in mean.index else (x0, y0)
            dx = int(round((x0 or mx) - mx))
            dy = int(round((y0 or my) - my))
            ws.cell(r, cdx, dx).fill = YELLOW
            ws.cell(r, cdy, dy).fill = YELLOW
            fx = f"=ROUND(AVERAGEIF({rng(K)},${SK}{r},{rng(X)})+{L(cdx)}{r},0)"
            fy = f"=ROUND(AVERAGEIF({rng(K)},${SK}{r},{rng(Y)})+{L(cdy)}{r},0)"
            ws.cell(r, h[f"{pre}_x_px"], fx).fill = GREY
            ws.cell(r, h[f"{pre}_y_px"], fy).fill = GREY
            ws.cell(r, h[f"{pre}_x_pct"], f"={L(h[pre + '_x_px'])}{r}/5120*100").fill = GREY
            ws.cell(r, h[f"{pre}_y_pct"], f"={L(h[pre + '_y_px'])}{r}/2880*100").fill = GREY
    if "_sync_state" in wb.sheetnames:
        del wb["_sync_state"]

    g = wb["안내"] if "안내" in wb.sheetnames else None
    if g is not None:
        g.append(["실시간 연동 (v2.6)",
                  "노란 칸만 고친다. Station_Visual_Nodes 의 x_px·y_px 를 고치면 역명·클릭 위치(회색, 수식)가 바로 따라온다. "
                  "역명만 옮기려면 Station_Labels 의 label_dx·label_dy(이동량)를 고친다. 다른 역과 좌표가 같은 행은 주황색으로 바로 표시된다. "
                  "미리보기: python scripts/v2/29_watch_map.py 를 켜 두면 저장할 때마다 브라우저 미리보기가 새로고침된다."])
    wb.save(BASE)
    kept = sum(1 for ws in (lab, wb["Station_Click_Areas"]) for r in range(2, ws.max_row + 1)
               for c in (ws.max_column - 1, ws.max_column) if ws.cell(r, c).value not in (0, None))
    print(f"변환 완료: {BASE}\n백업: {bak}\n직접 옮겨 둔 역명·클릭 위치 보존 (이동량이 0 이 아닌 칸 {kept}개)")
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from yeoyuro_v2.proc import utf8_stdout
    utf8_stdout()                       # Windows 에서 한글 출력 때문에 멈추지 않게
    raise SystemExit(main())
