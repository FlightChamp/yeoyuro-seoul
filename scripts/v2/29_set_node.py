"""
scripts/v2/29_set_node.py
=========================
엑셀을 열지 않고 기본 좌표 파일의 역 점(Station_Visual_Nodes) 좌표를 바꾼다.

    python scripts/v2/29_set_node.py 흑석_L9 2129 1750
    python scripts/v2/29_set_node.py 흑석_L9 2129 1750 봉은사_L9 2923 1964      # 여러 개 한 번에
    python scripts/v2/29_set_node.py 흑석                                       # 그 역의 노드 좌표만 보기

- 노드 이름은 '역명_L호선' (예: 동작_L4, 동작_L9). 역명만 쓰면 그 역의 노드 목록과 좌표를 보여 준다.
- 바꾼 뒤 python scripts/v2/29_check_map_workbook.py 를 실행하면 클릭 영역·역명이 같이 따라오고 미리보기가 만들어진다.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "data" / "master" / "yeoyuro_seoul_vector_map_coordinate_workbook.xlsx"


def main(argv=None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if not args:
        print(__doc__)
        return 1
    if (BASE.parent / ("~$" + BASE.name)).exists():
        print("[중단] 엑셀이 기본 파일을 열고 있습니다. 엑셀을 닫고 다시 실행하세요.")
        return 3
    wb = openpyxl.load_workbook(BASE)
    ws = wb["Station_Visual_Nodes"]
    h = {c.value: c.column for c in ws[1] if c.value}
    rows = {ws.cell(r, h["visual_node_id"]).value: r for r in range(2, ws.max_row + 1)}

    if len(args) == 1:                                   # 보기 모드
        key = args[0]
        hits = [n for n in rows if n == key or (n and n.rsplit("_L", 1)[0] == key)]
        if not hits:
            print(f"'{key}' 노드가 없습니다.")
            return 1
        for n in hits:
            r = rows[n]
            print(f"{n}: ({ws.cell(r, h['x_px']).value}, {ws.cell(r, h['y_px']).value})  {ws.cell(r, h['mark_status']).value}")
        return 0
    if len(args) % 3:
        print("형식: 노드이름 x y [노드이름 x y ...]")
        return 1
    done = []
    for i in range(0, len(args), 3):
        node, x, y = args[i], args[i + 1], args[i + 2]
        if node not in rows:
            print(f"[중단] '{node}' 노드가 없습니다. 역명만 넣어 노드 이름을 확인하세요 (예: python scripts/v2/29_set_node.py {node.split('_')[0]})")
            return 1
        x, y = int(float(x)), int(float(y))
        if not (0 <= x <= 5120 and 0 <= y <= 2880):
            print(f"[중단] {node}: 좌표가 캔버스(0~5120, 0~2880) 밖입니다.")
            return 1
        r = rows[node]
        old = (ws.cell(r, h["x_px"]).value, ws.cell(r, h["y_px"]).value)
        ws.cell(r, h["x_px"], x)
        ws.cell(r, h["y_px"], y)
        done.append(f"{node}: {old} → ({x}, {y})")
    wb.save(BASE)
    print("\n".join(done))
    print("다음: python scripts/v2/29_check_map_workbook.py  (클릭·역명 동기화 + 미리보기)")
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from yeoyuro_v2.proc import utf8_stdout
    utf8_stdout()                       # Windows 에서 한글 출력 때문에 멈추지 않게
    raise SystemExit(main())
